import json
from pathlib import Path
import sys

from pptx import Presentation
from pptx.inspect import BULLET_FOLLOWS_TEXT, effective_font, effective_paragraph_format, effective_shape_format


def bullet_value(paragraph):
    fmt = effective_paragraph_format(paragraph)
    if fmt.bullet.type == "none":
        return {"kind": "none", "marker": None, "font": None, "size_pt": None}
    font = effective_font(paragraph.runs[0])
    face = fmt.bullet_font.value
    if face == BULLET_FOLLOWS_TEXT:
        face = font.name.value
    size = fmt.bullet_size
    if size.value == BULLET_FOLLOWS_TEXT:
        points = font.size.value_pt
    elif size.value_pt is not None:
        points = size.value_pt
    elif isinstance(size.value, float) and font.size.value_pt is not None:
        points = round(size.value * font.size.value_pt, 4)
    else:
        points = None
    marker = fmt.bullet.char
    if fmt.bullet.type == "numbered":
        assert fmt.bullet.number_scheme == "arabicPeriod"
        marker = str(fmt.bullet.start_at) + "."
    return {"kind": fmt.bullet.type, "marker": marker, "font": face, "size_pt": points}


def solve(root):
    prs = Presentation(root / "eval_fixtures/pptx/account-template.pptx")
    rows, shapes = [], []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.name in {"Rollout banner", "Reference panel", "Review highlight"}:
                colors = effective_shape_format(shape)
                heading = next(s.text for s in slide.shapes if s.name == "Heading")
                shapes.append({"slide": heading, "object": shape.name,
                               "fill_rgb": colors.fill_rgb.value, "line_rgb": colors.line_rgb.value})
            if shape.name not in {"Priorities", "Template annotation"}:
                continue
            for i, paragraph in enumerate(shape.text_frame.paragraphs, 1):
                font = effective_font(paragraph.runs[0])
                rows.append({"slide": slide.shapes.title.text, "object": shape.name,
                             "paragraph": i, "font": font.name.value,
                             "size_pt": font.size.value_pt, "color_rgb": font.color_rgb.value,
                             "bullet": bullet_value(paragraph)})
    output = root / "evals/template-format-audit/report.json"
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps({"paragraphs": rows, "shapes": shapes}, indent=2, ensure_ascii=False) + "\n")


if __name__ == "__main__":
    solve(Path(sys.argv[1]) if len(sys.argv) > 1 else Path("/app"))
