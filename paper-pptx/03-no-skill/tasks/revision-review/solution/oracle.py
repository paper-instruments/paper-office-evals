import json
from pathlib import Path
import sys

from pptx import Presentation
from pptx.diff import diff_decks
from pptx.inspect import BULLET_FOLLOWS_TEXT, effective_font, effective_paragraph_format


def bullet_value(paragraph):
    fmt = effective_paragraph_format(paragraph)
    if fmt.bullet.type == "none":
        return {"kind": "none", "marker": None, "font": None, "size_pt": None}
    font = effective_font(paragraph.runs[0])
    face = fmt.bullet_font.value
    if face == BULLET_FOLLOWS_TEXT: face = font.name.value
    size = fmt.bullet_size
    if size.value == BULLET_FOLLOWS_TEXT: points = font.size.value_pt
    elif size.value_pt is not None: points = size.value_pt
    elif isinstance(size.value, float) and font.size.value_pt is not None: points = round(size.value * font.size.value_pt, 4)
    else: points = None
    marker = fmt.bullet.char if fmt.bullet.type != "numbered" else str(fmt.bullet.start_at) + "."
    return {"kind": fmt.bullet.type, "marker": marker, "font": face, "size_pt": points}


def title(slide):
    return next(shape.text for shape in slide.shapes if shape.name == "Heading")


def solve(root):
    fixtures = root / "eval_fixtures/pptx"
    results = []
    for pair in json.loads((fixtures / "reviews.json").read_text()):
        before, after = Presentation(fixtures / pair["before"]), Presentation(fixtures / pair["after"])
        comparison = diff_decks(before, after, detail="full")
        changes = []
        if comparison.slides_moved:
            changes.append({"kind": "slide_order", "slide": None, "shape_id": None,
                            "before": [title(s) for s in before.slides], "after": [title(s) for s in after.slides]})
        before_slides = {s.slide_id: s for s in before.slides}
        after_slides = {s.slide_id: s for s in after.slides}
        for delta in comparison.slide_changes:
            old, new = before_slides[delta.slide_id], after_slides[delta.slide_id]
            old_shapes = {s.shape_id: s for s in old.shapes}
            new_shapes = {s.shape_id: s for s in new.shapes}
            if delta.notes_change:
                changes.append({"kind": "notes", "slide": title(old), "shape_id": None,
                                "before": old.read_notes_text(), "after": new.read_notes_text()})
            for shape_id, old_shape in old_shapes.items():
                new_shape = new_shapes[shape_id]
                if old_shape.has_table:
                    a = [[c.text for c in row.cells] for row in old_shape.table.rows]
                    b = [[c.text for c in row.cells] for row in new_shape.table.rows]
                    kind = "table"
                elif old_shape.has_text_frame:
                    a = [p.text for p in old_shape.text_frame.paragraphs]
                    b = [p.text for p in new_shape.text_frame.paragraphs]
                    kind = "text"
                else:
                    continue
                if a != b:
                    changes.append({"kind": kind, "slide": title(old), "shape_id": shape_id, "before": a, "after": b})
                elif old_shape.has_text_frame:
                    a = [bullet_value(p) for p in old_shape.text_frame.paragraphs]
                    b = [bullet_value(p) for p in new_shape.text_frame.paragraphs]
                    if a != b:
                        changes.append({"kind": "list", "slide": title(old), "shape_id": shape_id, "before": a, "after": b})
        results.append({"review": pair["review"], "changes": changes})
    output = root / "evals/revision-review/report.json"
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps({"reviews": results}, indent=2, ensure_ascii=False) + "\n")


if __name__ == "__main__":
    solve(Path(sys.argv[1]) if len(sys.argv) > 1 else Path("/app"))
